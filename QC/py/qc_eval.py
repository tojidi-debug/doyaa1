"""[품감] 품질관리수준 평가표(엑셀) — 심의안 초안 화면의 체크 상태로 만든다 (2026-10-08).

서식(docs/qc 「품질관리수준평가표.xlsx」)의 칸 배치·배점·셈법을 그대로 따르되 **새로 그린다**:
  - 점검회사 열은 회사 수만큼(2개사·3개사·4개사… 제한 없음). 서식은 4열 고정이고 #REF! 수식이 섞여 있어 그대로 채우지 않는다.
  - 합계·평균·점수는 엑셀 수식으로 두어(SUM·IF) 담당자가 엑셀에서 숫자를 고치면 다시 셈해진다. 값은 frontend/qc/qc_draft.js `score()` 와 같은 셈법.
셈법(서식 「배점 산정 근거」):
  품질관리절차 60점 = 감사시간 10(5:3) + 사전심리 20(5:5:3) + 조서 20(5:5:5) + 모니터링 10(5:3); 항목 점수 = 미설계·미운영 0 · 일부미흡 ½ · 지적없음 전부
  개별감사업무 40점 = 항목 배부점수(5·5·1·5·5·1·1·5·1·15·3·3·3·1·10·1, 합 65) × 40/65; 항목 점수 = 배부점수 × (1 − 평균지적개수),
    ②·③은 합쳐서 (1 − 평균/2) · 기타 주요 계정과목: 평균 1 초과 2 이하 ⅔ · 2 초과 4 이하 ⅓ · 4 초과 0 · 기타 감사조서의 문서화: 1 초과 2 이하 ½ · 2 초과 0 · 1 이하 전부
  법규위반 차감 = Σ(건수 × 차감점수 5/3/1) · 합계 90 이상 양호 · 50 이상 보통 · 그 밖 미흡
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ORG_ROWS = [  # (묶음, 항목 키, 항목 이름, 배부점수, 묶음 환산점수)
    ("감사조서 작성∙관리절차", "a1", "① 조서관리담당자 지정과 최종 감사파일 취합", 5, 20),
    ("감사조서 작성∙관리절차", "a2", "② 감사조서 보존∙관리", 5, 20),
    ("감사조서 작성∙관리절차", "a3", "③ 감사조서 관리대장 작성과 유지(입출고, 폐기)", 5, 20),
    ("업무품질관리검토(사전심리)절차", "b1", "① 업무품질관리검토절차 수립과 업무배정", 5, 20),
    ("업무품질관리검토(사전심리)절차", "b2", "② 사전심리 수행과 해결사항의 문서화", 5, 20),
    ("업무품질관리검토(사전심리)절차", "b3", "③ 감사보고서 발행 통제", 3, 20),
    ("감사시간 관리절차", "c1", "① 감사시간 집계·검증시스템 구축·운영", 5, 10),
    ("감사시간 관리절차", "c2", "② 감사시간 모니터링", 3, 10),
    ("모니터링(사후심리)", "d1", "① 사후심리 관련 정책과 절차", 5, 10),
    ("모니터링(사후심리)", "d2", "② 사후심리결과에 대한 관리", 3, 10),
]
INDI_ROWS = [  # (부문, 점수 키, 세는 항목 키들, 이름, 배부점수)
    ("감사계획의 수립과 평가절차", "i1", ["i1"], "① 감사계획 수립절차", 5),
    ("감사계획의 수립과 평가절차", "i23", ["i2", "i3"], "② 중요한 왜곡표시위험의 식별과 평가절차/평가된 왜곡표시위험 관련 대응절차", 5),
    ("감사계획의 수립과 평가절차", "i4", ["i4"], "③ 내부회계관리제도 검토절차", 1),
    ("감사증거의 수집과 문서화", "i5", ["i5"], "④ 재고자산 실재성과 상태 관련 감사절차", 5),
    ("감사증거의 수집과 문서화", "i6", ["i6"], "⑤ 외부조회 관련 감사절차", 5),
    ("감사증거의 수집과 문서화", "i7", ["i7"], "⑥ 표본감사 관련 감사절차", 1),
    ("감사증거의 수집과 문서화", "i8", ["i8"], "⑦ 초도감사시 기초잔액에 대한 감사절차", 1),
    ("감사증거의 수집과 문서화", "i9", ["i9"], "⑧ 특수관계자 거래에 대한 감사절차", 5),
    ("감사증거의 수집과 문서화", "i10", ["i10"], "⑨ 계속기업 존속능력에 대한 감사절차", 1),
    ("감사증거의 수집과 문서화", "i11", ["i11"], "⑩ 기타 주요 계정과목에 대한 감사절차", 15),
    ("감사의 종료절차", "i12", ["i12"], "⑪ 전반적인 결론을 위한 분석적절차", 3),
    ("감사의 종료절차", "i13", ["i13"], "⑫ 후속사건 감사절차", 3),
    ("감사의 종료절차", "i14", ["i14"], "⑬ 서면진술 관련 문서화", 3),
    ("감사의 종료절차", "i15", ["i15"], "⑭ 지배기구와의 커뮤니케이션절차", 1),
    ("감사의 종료절차", "i16", ["i16"], "⑮ 기타 감사조서의 문서화", 10),
    ("그룹재무제표 감사절차", "i17", ["i17"], "⑯ 그룹재무제표 감사절차(감사계획 수립·독립성 확인)", 1),
]
INDI_SUM = sum(r[4] for r in INDI_ROWS)   # 65
VIOL_ROWS = ["GAAP", "GAAS", "품질관리전담자", "연속감사제한", "미등기이사 업무수행", "기타"]

HEAD = PatternFill("solid", fgColor="FFF2CC")
SUB = PatternFill("solid", fgColor="FFFBE6")
INPUT = PatternFill("solid", fgColor="E8F4FF")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F = Font(name="맑은 고딕", size=10)
FB = Font(name="맑은 고딕", size=10, bold=True)
FT = Font(name="맑은 고딕", size=14, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _put(ws, row: int, col: int, value: Any, *, font=F, fill=None, align=CENTER, border=True, fmt: str | None = None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font; c.alignment = align
    if fill is not None:
        c.fill = fill
    if border:
        c.border = BOX
    if fmt:
        c.number_format = fmt
    return c


def build_xlsx(payload: dict[str, Any]) -> bytes:
    """payload = {firm, period, prev_period, cpa:[당,전], audited:[당,전], fee:[당,전], companies:[이름…],
                  org:{a1:'일부미흡'|'미운영'|'미설계'|'미지적'…}, indi:{i1:[회사별 수…]…}, viol:{GAAP:{n,pts}…}, level_reason, score:{…}}"""
    wb = Workbook(); ws = wb.active; ws.title = "수준평가"
    companies = [str(c or f"회사{i + 1}") for i, c in enumerate(payload.get("companies") or [])]
    K = len(companies)
    for col, w in zip("ABCDEFGHIJ", (2, 26, 44, 12, 12, 12, 12, 12, 12, 12)):
        ws.column_dimensions[col].width = w
    for i in range(K):
        ws.column_dimensions[get_column_letter(4 + i)].width = 12
    r = 1
    _put(ws, r, 2, "회계법인 품질관리수준 평가", font=FT, align=LEFT, border=False); r += 1
    _put(ws, r, 2, "평가목적 : 감사인 감리 실시결과 회계법인의 '품질관리수준'을 평가하여 수준별(양호, 보통, 미흡)로 인센티브/패널티 부과 및 감사인 지정을 위한 기초자료로 활용", align=LEFT, border=False); r += 2

    # 1. 감사인 개요
    _put(ws, r, 2, "1. 감사인 개요", font=FB, align=LEFT, border=False); r += 1
    _put(ws, r, 2, "ㅇ 회계법인명", align=LEFT, border=False); _put(ws, r, 3, f": {payload.get('firm', '')}", align=LEFT, border=False); r += 1
    _put(ws, r, 2, "ㅇ 외부감사현황", align=LEFT, border=False); _put(ws, r, 5, "(단위: 사, 명, 백만원)", align=LEFT, border=False); r += 1
    _put(ws, r, 2, "사 업 연 도", fill=HEAD); _put(ws, r, 3, "", fill=HEAD); _put(ws, r, 4, payload.get("period", ""), fill=HEAD); _put(ws, r, 5, f"{payload.get('prev_period', '')}*", fill=HEAD); r += 1
    first = r
    for label, key in (("공인회계사 수", "cpa"), ("외부감사회사 수", "audited"), ("외부감사수임료 합계", "fee")):
        v = payload.get(key) or [None, None]
        _put(ws, r, 2, label, align=LEFT); _put(ws, r, 3, ""); _put(ws, r, 4, v[0], fmt="#,##0"); _put(ws, r, 5, v[1], fmt="#,##0"); r += 1
    _put(ws, r, 2, "공인회계사 1인당 외부감사 수", align=LEFT); _put(ws, r, 3, "")
    _put(ws, r, 4, f"=IF(D{first}=0,\"\",D{first + 1}/D{first})", fmt="0.0"); _put(ws, r, 5, f"=IF(E{first}=0,\"\",E{first + 1}/E{first})", fmt="0.0"); r += 1
    _put(ws, r, 2, "외부감사 평균 수임료", align=LEFT); _put(ws, r, 3, "")
    _put(ws, r, 4, f"=IF(D{first + 1}=0,\"\",D{first + 2}/D{first + 1})", fmt="0.0"); _put(ws, r, 5, f"=IF(E{first + 1}=0,\"\",E{first + 2}/E{first + 1})", fmt="0.0"); r += 1
    _put(ws, r, 2, "* 전기 외부감사현황은 정보값을 입수가능한 경우에 입력", align=LEFT, border=False); r += 2

    # 2. 평가
    _put(ws, r, 2, "2. 개선권고사항에 따른 품질관리수준 평가", font=FB, align=LEFT, border=False); r += 1
    _put(ws, r, 2, "(산정기준) 1. '품질관리절차'와 '개별감사업무'의 중요도 가중치 = 60% : 40%  2. '품질관리절차'는 기존 점검 항목(감사조서 작성∙관리절차와 사전심리절차)의 비중과 신규 점검항목(감사시간 관리·모니터링(사후심리))의 비중은 2:1", align=LEFT, border=False); r += 2
    # 품질관리절차
    for c, t in zip(range(2, 10), ("구성요소", "주요 점검 항목", "미운영", "일부미흡", "지적없음", "개선권고유형", "배부점수(환산)", "점수")):
        _put(ws, r, c, t, font=FB, fill=HEAD)
    r += 1
    org = payload.get("org") or {}
    org_first = r
    for grp, key, name, pts, gsum in ORG_ROWS:
        cls = org.get(key) or "미지적"
        gw = sum(x[3] for x in ORG_ROWS if x[0] == grp)
        w = gsum * pts / gw
        _put(ws, r, 2, grp, fill=SUB, align=LEFT); _put(ws, r, 3, name, align=LEFT)
        _put(ws, r, 4, "O" if cls in ("미운영", "미설계") else "", fill=INPUT); _put(ws, r, 5, "O" if cls == "일부미흡" else "", fill=INPUT)
        _put(ws, r, 6, f'=IF(OR(D{r}="O",E{r}="O"),"","O")')
        _put(ws, r, 7, f'=IF(D{r}="O","미운영",IF(E{r}="O","일부미흡","지적없음"))')
        _put(ws, r, 8, round(w, 4), fmt="0.00")
        _put(ws, r, 9, f'=IF(D{r}="O",0,IF(E{r}="O",H{r}/2,H{r}))', fmt="0.00")
        r += 1
    org_last = r - 1
    _put(ws, r, 2, "품질관리절차 관련 점수(60점 만점)", font=FB, fill=HEAD, align=LEFT);
    for c in range(3, 9): _put(ws, r, c, "", fill=HEAD)
    _put(ws, r, 9, f"=SUM(I{org_first}:I{org_last})", font=FB, fill=HEAD, fmt="0.00"); org_total_row = r; r += 2

    # 개별감사업무 — 서식과 같은 줄: ②는 「1. 식별과 평가」·「2. 대응」의 합, ⑮는 하위 5줄의 합(수식). 점수는 상위 줄에서(사용자 지정 2026-10-08)
    co_cols = list(range(4, 4 + K))
    avg_col = 4 + K; w_col = avg_col + 1; pts_col = avg_col + 2
    _put(ws, r, 2, "개별감사업무 점검회사 수", font=FB, fill=HEAD, align=LEFT); _put(ws, r, 3, K, fill=INPUT)
    kcell = f"$C${r}"
    _put(ws, r, 4, "← 평균지적개수 = 회사별 지적 수 합 ÷ 회사 수", align=LEFT, border=False); r += 1
    _put(ws, r, 2, "구성요소", font=FB, fill=HEAD); _put(ws, r, 3, "주요 점검 항목", font=FB, fill=HEAD)
    for i, name in enumerate(companies):
        _put(ws, r, 4 + i, name, font=FB, fill=HEAD)
    _put(ws, r, avg_col, "평균지적개수", font=FB, fill=HEAD); _put(ws, r, w_col, "배부점수(환산)", font=FB, fill=HEAD); _put(ws, r, pts_col, "점수", font=FB, fill=HEAD)
    r += 1
    indi = payload.get("indi") or {}
    subs_in = payload.get("indi_subs") or {}
    SUB_LABELS = {
        "i11": [("s1", "대손충당금 적정성(매출채권 연령분석 등) 검토 미흡"), ("s2", "재고자산 평가(저가법·진부화 등) 검토 미흡"), ("s3", "수익인식(매출) 감사절차 미흡"), ("s4", "지분법 회계처리 검토 미흡"), ("s5", "기타")],
        "i16": [("s1", "공시사항점검표 작성 미흡"), ("s2", "완결조서/감사종료 점검서식 작성 미흡"), ("s3", "수행자/검토자 업무수행일자/서명미흡"), ("s4", "조서간 유기적 연결 미흡"), ("s5", "기타")],
        "i17": [("s1", "그룹감사업무 수행시 감사계획 수립절차"), ("s2", "그룹감사업무 수행시 독립성 확인절차")],   # 상위 = 둘 중 하나라도 있으면 1
    }
    val = lambda key, i: int((indi.get(key) or [])[i]) if i < len(indi.get(key) or []) else 0
    score_rows: list[int] = []

    def points(rr: int, skey: str, G: float) -> None:
        rng = f"{get_column_letter(4)}{rr}:{get_column_letter(3 + K)}{rr}" if K else None
        _put(ws, rr, avg_col, f"=IF({kcell}=0,0,SUM({rng})/{kcell})" if rng else 0, fmt="0.00"); _put(ws, rr, w_col, round(G, 4), fmt="0.00")
        A, W = f"{get_column_letter(avg_col)}{rr}", f"{get_column_letter(w_col)}{rr}"
        if skey == "i23":
            f = f"={W}-{W}*({A}/2)"
        elif skey == "i11":
            f = f"=IF({A}>4,0,IF({A}>2,{W}/3,IF({A}>1,{W}*2/3,{W}-{W}*{A})))"
        elif skey == "i16":
            f = f"=IF({A}>2,0,IF({A}>1,{W}/2,{W}))"
        else:
            f = f"={W}-{W}*{A}"
        _put(ws, rr, pts_col, f, fmt="0.00"); score_rows.append(rr)

    def blanks(rr: int) -> None:
        for c in (avg_col, w_col, pts_col):
            _put(ws, rr, c, "", fill=SUB)

    indi_first = r
    for part, skey, keys, name, pts in INDI_ROWS:
        G = 40 * pts / INDI_SUM
        if skey == "i23":
            parent = r
            _put(ws, r, 2, part, fill=SUB, align=LEFT); _put(ws, r, 3, name, align=LEFT, font=FB)
            for i in range(K):
                col = get_column_letter(4 + i); _put(ws, r, 4 + i, f"=SUM({col}{r + 1}:{col}{r + 2})")
            points(r, skey, G); r += 1
            for no, (k2, label) in enumerate((("i2", "중요한 왜곡표시위험의 식별과 평가절차"), ("i3", "평가된 왜곡표시위험 관련 대응절차")), 1):
                _put(ws, r, 2, part, fill=SUB, align=LEFT); _put(ws, r, 3, f"  {no}. {label}", align=LEFT)
                for i in range(K):
                    _put(ws, r, 4 + i, val(k2, i), fill=INPUT)
                blanks(r); r += 1
            continue
        if skey in SUB_LABELS:   # ⑩·⑮ — 상위 줄은 하위 줄의 SUM, 점수는 상위 줄(사용자 지정 2026-10-08)
            labels = SUB_LABELS[skey]; sub = subs_in.get(skey) or {}
            _put(ws, r, 2, part, fill=SUB, align=LEFT); _put(ws, r, 3, name, align=LEFT, font=FB)
            for i in range(K):
                col = get_column_letter(4 + i); rng = f"{col}{r + 1}:{col}{r + len(labels)}"
                _put(ws, r, 4 + i, f"=IF(SUM({rng})>0,1,0)" if skey == "i17" else f"=SUM({rng})")
            points(r, skey, G); r += 1
            sub_vals = {sk: [int((sub.get(sk) or [])[i]) if i < len(sub.get(sk) or []) else 0 for i in range(K)] for sk, _ in labels}
            for i in range(K):   # 지적 예시로만 고른 것(하위 줄 체크 없이)은 「기타」에 넣어 상위 합계를 화면 건수와 맞춘다
                rest = val(skey, i) - sum(sub_vals[sk][i] for sk, _ in labels)
                if rest > 0:
                    sub_vals["s5" if "s5" in sub_vals else "s1"][i] += rest
            for no, (sk, label) in enumerate(labels, 1):
                _put(ws, r, 2, part, fill=SUB, align=LEFT); _put(ws, r, 3, f"  {no}. {label}", align=LEFT)
                for i in range(K):
                    _put(ws, r, 4 + i, sub_vals[sk][i], fill=INPUT)
                blanks(r); r += 1
            continue
        _put(ws, r, 2, part, fill=SUB, align=LEFT); _put(ws, r, 3, name, align=LEFT)
        for i in range(K):
            _put(ws, r, 4 + i, sum(val(k2, i) for k2 in keys), fill=INPUT)
        points(r, skey, G); r += 1
    _put(ws, r, 2, "개별감사업무 관련 점수(40점 만점)", font=FB, fill=HEAD, align=LEFT)
    for c in range(3, pts_col): _put(ws, r, c, "", fill=HEAD)
    P0 = get_column_letter(pts_col)
    _put(ws, r, pts_col, "=" + "+".join(f"{P0}{x}" for x in score_rows), font=FB, fill=HEAD, fmt="0.00"); indi_total_row = r; r += 2

    # 법규위반 차감
    _put(ws, r, 2, "법규위반 차감", font=FB, fill=HEAD, align=LEFT); _put(ws, r, 3, "구 분", font=FB, fill=HEAD); _put(ws, r, 4, "건 수", font=FB, fill=HEAD); _put(ws, r, 5, "차감점수(5/3/1)", font=FB, fill=HEAD); _put(ws, r, 6, "차감 계", font=FB, fill=HEAD); r += 1
    viol = payload.get("viol") or {}
    v_first = r
    for name in VIOL_ROWS:
        v = viol.get(name) or {}
        _put(ws, r, 2, ""); _put(ws, r, 3, name, align=LEFT); _put(ws, r, 4, int(v.get("n") or 0), fill=INPUT); _put(ws, r, 5, int(v.get("pts") or 0), fill=INPUT); _put(ws, r, 6, f"=D{r}*E{r}"); r += 1
    v_last = r - 1
    _put(ws, r, 2, ""); _put(ws, r, 3, "법규 위반 차감 계", font=FB, align=LEFT); _put(ws, r, 4, f"=SUM(D{v_first}:D{v_last})"); _put(ws, r, 5, ""); _put(ws, r, 6, f"=SUM(F{v_first}:F{v_last})", font=FB); v_total_row = r; r += 1
    _put(ws, r, 3, "* 위반자별 5점 차감(예: 미등기이사의 감사업무 수행, 위반자 2명 → 10점). 손해배상준비금 미적립처럼 품질관리제도의 영향이 없는 사항은 반영하지 않음", align=LEFT, border=False); r += 2

    # 합계
    P = get_column_letter(pts_col)
    _put(ws, r, 2, "품질관리점수 합계", font=FB, fill=HEAD, align=LEFT); _put(ws, r, 3, f"=I{org_total_row}+{P}{indi_total_row}-F{v_total_row}", font=FB, fmt="0.00"); total_row = r; r += 1
    _put(ws, r, 2, "품질관리수준", font=FB, fill=HEAD, align=LEFT); _put(ws, r, 3, f'=IF(C{total_row}>=90,"양호",IF(C{total_row}>=50,"보통","미흡"))', font=FB); r += 1
    _put(ws, r, 2, "* 점수에 따른 품질관리수준 — 양호 90점 이상 ~ 100점 · 보통 50점 이상 ~ 90점 미만 · 미흡 50점 미만", align=LEFT, border=False); r += 2

    # 3·4
    _put(ws, r, 2, "3. 기타 고려요소 (점수에 따른 품질관리수준이 적합하지 않은 경우 기술)", font=FB, align=LEFT, border=False); r += 1
    _put(ws, r, 2, "감사인 특성(감사업무 비중, 구성원 특성, 실질적 감사팀 운영 등), 질적 고려 요소(직전 지적사항의 개선정도, 운영규약의 강제성, 실질적 감사보수 배분, 조서 아카이브 준수, 법규위반 내용 등) 판단이유에 기재", align=LEFT, border=False); r += 2
    _put(ws, r, 2, "4. 최종 품질관리수준 평가", font=FB, align=LEFT, border=False); r += 1
    _put(ws, r, 2, "최종 수준평가", font=FB, fill=HEAD); _put(ws, r, 3, "판 단 이 유", font=FB, fill=HEAD); r += 1
    _put(ws, r, 2, payload.get("level") or f"=C{total_row + 1}", fill=INPUT); c = _put(ws, r, 3, payload.get("level_reason") or "", fill=INPUT, align=LEFT); ws.row_dimensions[r].height = 60; r += 2

    # 지적사항 요약(회사별)
    _put(ws, r, 2, "<개별감사업무 관련 지적사항 — 회사별 지적 수>", font=FB, align=LEFT, border=False); r += 1
    _put(ws, r, 2, "구성요소", font=FB, fill=HEAD); _put(ws, r, 3, "주요 점검항목", font=FB, fill=HEAD)
    for i, name in enumerate(companies):
        _put(ws, r, 4 + i, name, font=FB, fill=HEAD)
    _put(ws, r, 4 + K, "계", font=FB, fill=HEAD); _put(ws, r, 5 + K, "지적 회사 수", font=FB, fill=HEAD); r += 1
    s_first = r
    for part, skey, keys, name, pts in INDI_ROWS:
        _put(ws, r, 2, part, fill=SUB, align=LEFT); _put(ws, r, 3, name, align=LEFT)
        for i in range(K):
            n = sum(int((indi.get(k) or [0] * K)[i] if i < len(indi.get(k) or []) else 0) for k in keys)
            _put(ws, r, 4 + i, n)
        if K:
            rng = f"{get_column_letter(4)}{r}:{get_column_letter(3 + K)}{r}"
            _put(ws, r, 4 + K, f"=SUM({rng})"); _put(ws, r, 5 + K, f"=COUNTIF({rng},\">=1\")")
        else:
            _put(ws, r, 4 + K, 0); _put(ws, r, 5 + K, 0)
        r += 1
    _put(ws, r, 2, "합계", font=FB, fill=HEAD); _put(ws, r, 3, "", fill=HEAD)
    for i in range(K):
        col = get_column_letter(4 + i); _put(ws, r, 4 + i, f"=SUM({col}{s_first}:{col}{r - 1})", font=FB, fill=HEAD)
    for c in (4 + K, 5 + K):
        col = get_column_letter(c); _put(ws, r, c, f"=SUM({col}{s_first}:{col}{r - 1})", font=FB, fill=HEAD)
    r += 2
    _put(ws, r, 2, "※ 기재요령: 항목별 지적사항이 있으면 1, 없으면 0. 기타 주요 계정과목에 대한 감사절차와 기타 감사조서의 문서화 항목은 세부 지적건수. 품질관리절차는 미설계·미운영 / 일부미흡 칸에 'O'.", align=LEFT, border=False)
    ws.freeze_panes = "A5"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
