"""[브라우저(Pyodide)판 — 자동 생성, 손으로 고치지 말 것. 원본: backend/src/gumiho/services/qc_audited.py, 전자공시 검색만 구글 Apps Script 중계가 함]
외감회사대사 — 회계법인이 최근 1년 동안 전자공시에 제출한 감사보고서의 회사 목록(공시 기준)과
회계법인이 제출한 피외감대상회사 목록을 맞춰 본다 ([품감] 심의안(회계법인), 2026-10-08).

    공시 목록   dart.fss.or.kr 공시통합검색(dsab007) 을 **제출인명**으로 검색한다. OpenDART list.json 은 제출인으로 못 거른다
               (회사 고유번호 기준이라, 한 해 외부감사 공시 전부를 받아 거르면 수백 번 부른다). 검색은 로그인·키 없이 열린다.
               제출인명은 공시 표기대로 **빈칸 없이**(「회계법인나루」). 보고서명 「감사보고서 (YYYY.MM)」·「연결감사보고서 (YYYY.MM)」만
               쓰고, 정정본은 접수일이 늦은 것 하나로 합친다(연도 = 보고서명의 YYYY).
    엑셀       연도마다 시트: 머리줄(개별 N사 · 연결 M사), A 회사명 · B 감사보고서 주소 · C 연결감사보고서 주소(하이퍼링크).
    대사       회사명을 씻어(주식회사·유한회사·농업회사법인 … 뗌, 빈칸·구두점 뗌, 제이호↔제2호, 스펙↔SPEC) 맞춘다.
               안 맞는 것은 비슷한 이름(바이그램 Dice ≥ .8)을 후보로 보여 준다. 짐작해서 맞췄다고 하지 않는다 — 「확인 필요」로 둔다.
"""
from __future__ import annotations

import html
import io
import logging
import re
import unicodedata
from datetime import date, timedelta
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

DART = "https://dart.fss.or.kr"
HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
           "Accept-Language": "ko-KR,ko;q=0.9"}
TIMEOUT = 30.0
MAX_PAGES = 30
REPORT_RE = re.compile(r"(\[[^\]]*\]\s*)?(연결)?감사보고서\s*\((\d{4})\.(\d{2})\)")


# ── 1. 공시 목록 ────────────────────────────────────────────────────────────

def _cells(row_html: str) -> list[str]:
    return [re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", c))).strip() for c in re.findall(r"<td[^>]*>(.*?)</td>", row_html, re.S)]


def _company(row_html: str) -> str:
    """회사명 칸의 링크 글만(앞의 시장 표시 「기」·「유」 같은 배지는 뺀다)."""
    tds = re.findall(r"<td[^>]*>(.*?)</td>", row_html, re.S)
    if len(tds) < 2:
        return ""
    m = re.search(r"<a[^>]*>(.*?)</a>", tds[1], re.S)
    text = m.group(1) if m else tds[1]
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", text))).strip()


def parse_search(page_html: str) -> tuple[list[dict[str, Any]], int, int]:
    """검색 결과 한 쪽 → (행 목록, 지금 쪽, 전체 쪽). 행: {company, report_nm, filer, rcept_dt, rcept_no}"""
    rows = []
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", page_html, re.S):
        cells = _cells(tr)
        no = re.search(r"rcpNo=(\d{14})", tr)
        if len(cells) < 5 or not no:
            continue
        rows.append({"company": _company(tr), "report_nm": cells[2], "filer": cells[3], "rcept_dt": cells[4], "rcept_no": no.group(1)})
    m = re.search(r"\[(\d+)/(\d+)\]", page_html)
    cur, total = (int(m.group(1)), int(m.group(2))) if m else (1, 1)
    return rows, cur, total


def default_min_period(max_period: str, today: date | None = None) -> str:
    """집계 하한 결산기 = 결산기 상한 연도의 전년 1월(2026.03 → 2025.01: 최근 사업연도 전부 + 상한까지 석 달). 상한이 없으면 올해의 전년 1월."""
    y = int(max_period[:4]) if re.fullmatch(r"\d{4}\.\d{2}", max_period or "") else (today or date.today()).year
    return f"{y - 1}.01"


def group_audits(rows: list[dict[str, Any]], max_period: str = "", min_period: str = "") -> dict[str, Any]:
    """감사보고서·연결감사보고서를 (연도, 회사) 로 묶는다. 정정본은 접수일이 늦은 것.
    `max_period`(YYYY.MM, 회계법인 사업연도 말)를 주면 그보다 뒤 결산의 보고서(예: 2026.04)는 뺀다(사용자 지정 2026-10-08).
    `min_period` 보다 앞선 결산(과거 연도의 정정 등)은 **집계·대사에서 빼되**(old=True) 엑셀 연도 시트에는 남긴다(사용자 지정 2026-10-08)."""
    years: dict[int, dict[str, dict[str, Any]]] = {}
    skipped = 0
    for r in rows:
        m = REPORT_RE.search(r["report_nm"])
        if not m or not r["company"]:
            skipped += 1
            continue
        if max_period and f"{m.group(3)}.{m.group(4)}" > max_period:
            skipped += 1
            continue
        year, kind = int(m.group(3)), ("con" if m.group(2) else "indi")
        period = f"{m.group(3)}.{m.group(4)}"
        key = norm_name(r["company"]) + "|" + period   # 같은 회사라도 결산기가 다르면 보고서가 둘(사업보고서 감사실적도 보고서 수로 센다)
        ent = years.setdefault(year, {}).setdefault(key, {"company": r["company"], "indi": None, "con": None, "indi_dt": "", "con_dt": "", "period": period, "corrected": False, "rows": [],
                                                           "old": bool(min_period) and period < min_period})
        url = f"{DART}/dsaf001/main.do?rcpNo={r['rcept_no']}"
        ent["rows"].append({"report_nm": r["report_nm"], "rcept_dt": r["rcept_dt"]})
        if "정정" in r["report_nm"]:
            ent["corrected"] = True
        if r["rcept_dt"] >= ent[f"{kind}_dt"]:
            ent[kind], ent[f"{kind}_dt"] = url, r["rcept_dt"]
    out_years = []
    for y in sorted(years, reverse=True):
        items = sorted(years[y].values(), key=lambda e: (e["company"], e["period"]))
        live = [e for e in items if not e["old"]]
        out_years.append({"year": y, "items": items, "indi_count": sum(1 for e in live if e["indi"]), "con_count": sum(1 for e in live if e["con"]),
                          "old_count": sum(1 for e in items if e["old"] and e["indi"])})
    live_cos = {norm_name(e["company"]) for y in years.values() for e in y.values() if e["indi"] and not e["old"]}
    return {"years": out_years, "total_rows": len(rows), "skipped": skipped, "min_period": min_period, "indi_companies": len(live_cos)}


# ── 2. 이름 씻기 ────────────────────────────────────────────────────────────

_LEGAL = re.compile(r"주식회사|유한회사|유한책임회사|합자회사|합명회사|농업회사법인|어업회사법인|영농조합법인|영어조합법인|사단법인|재단법인|학교법인|의료법인|사회복지법인|협동조합|㈜|\(주\)|\(유\)|\(합\)|\(사\)|\(재\)|co\.?,?\s*ltd\.?|corp(?:oration)?\.?|inc\.?|ltd\.?|limited|company", re.I)
_NOISE = re.compile(r"[\s\-·ㆍ・,.'\"()\[\]{}（）「」『』_/&+!?:;]")
_KNUM = {"영": 0, "일": 1, "이": 2, "삼": 3, "사": 4, "오": 5, "육": 6, "칠": 7, "팔": 8, "구": 9, "십": 10}
#: 같은 뜻의 다른 표기(사용자 확인 2026-10-08): 리츠 ↔ (위탁관리)부동산투자회사 · 앤/엔/앤드/& · 피에브이(공시 오타) ↔ 피에프브이
_ALIAS = [("스펙", "spec"), ("스팩", "spac"), ("에스피에이씨", "spac"), ("에이비에스", "abs"), ("제일차", "1차"), ("제이차", "2차"),
          ("기업구조조정부동산투자회사", "리츠"), ("위탁관리부동산투자회사", "리츠"), ("자기관리부동산투자회사", "리츠"), ("부동산투자회사", "리츠"), ("reits", "리츠"), ("reit", "리츠"),
          ("피에프브이", "pfv"), ("피에브이", "pfv"), ("앤드", "앤"), ("엔", "앤"), ("&", "앤")]
#: 괄호 안의 옛 이름·주석 「(구, 에스에스개발(주))」·「(舊 …)」 은 뗀다
_FORMER = re.compile(r"\(\s*(?:구\s*[,，:]|구\s+|舊|전\s*[,，:]|前)[^()]*(?:\([^()]*\))?[^()]*\)")   # 「( 구 :주식회사 부광테크)」처럼 괄호 뒤 빈칸도
#: 영문자 이름을 한글로 적은 것(케이비 → kb · 에스케이 → sk). 두 글자 이상 이어질 때만 바꾼다
_LETTERS = [("더블유", "w"), ("에이치", "h"), ("에이", "a"), ("비", "b"), ("씨", "c"), ("디", "d"), ("이", "e"), ("에프", "f"), ("지", "g"), ("아이", "i"),
            ("제이", "j"), ("케이", "k"), ("엘", "l"), ("엠", "m"), ("엔", "n"), ("오", "o"), ("피", "p"), ("큐", "q"), ("알", "r"), ("에스", "s"),
            ("티", "t"), ("유", "u"), ("브이", "v"), ("엑스", "x"), ("와이", "y"), ("제트", "z")]
_LETTER_RE = re.compile("(?:" + "|".join(re.escape(k) for k, _ in sorted(_LETTERS, key=lambda x: -len(x[0]))) + "){2,}")
_LETTER_MAP = dict(_LETTERS)


def _letters(s: str) -> str:
    def conv(m: re.Match) -> str:
        chunk = m.group(0); out = ""
        while chunk:
            for k in sorted(_LETTER_MAP, key=len, reverse=True):
                if chunk.startswith(k):
                    out += _LETTER_MAP[k]; chunk = chunk[len(k):]; break
            else:
                return m.group(0)
        return out
    return _LETTER_RE.sub(conv, s)


def _knum(s: str) -> str:
    """제이호 → 제2호 · 제십일호 → 제11호 (제N호·제N차 꼴만)."""
    def conv(m: re.Match) -> str:
        word = m.group(2)
        n, cur = 0, 0
        for ch in word:
            v = _KNUM.get(ch)
            if v is None:
                return m.group(0)
            if v == 10:
                n += (cur or 1) * 10
                cur = 0
            else:
                cur = v
        return f"{m.group(1)}{n + cur}{m.group(3)}"
    s = re.sub(r"(제)([영일이삼사오육칠팔구십]{1,4})(호|차|기|회)", conv, s)
    return re.sub(r"(?<![영일이삼사오육칠팔구십])()([영일이삼사오육칠팔구십]{1,3})(피에프브이|pfv|호|차|단지)", conv, s)


def norm_name(value: str, aliases: dict[str, str] | None = None) -> str:
    s = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    s = _FORMER.sub("", s)
    s = _LEGAL.sub("", s)
    s = _NOISE.sub("", s)
    s = _knum(s)
    for a, b in _ALIAS:
        s = s.replace(a, b)
    s = _letters(s)
    if aliases and s in aliases:   # 사용자가 「같은 회사」로 등록한 짝
        s = aliases[s]
    return s


def plain_name(value: str) -> str:
    """법인격 표시(주식회사·㈜·(주)·유한회사 …)와 빈칸·대소문자만 뗀 이름 — 「표기 다름」은 이것이 다를 때만(제이호↔제2호, 피에프브이↔PFV 같은 것)."""
    s = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    s = _LEGAL.sub("", s)
    return re.sub(r"[\s\-·ㆍ・,.'\"()\[\]{}（）]", "", s)


def _bigrams(s: str) -> dict[str, int]:
    out: dict[str, int] = {}
    for k in range(len(s) - 1):
        out[s[k:k + 2]] = out.get(s[k:k + 2], 0) + 1
    return out


def similarity(a: str, b: str) -> float:
    A, B = _bigrams(a), _bigrams(b)
    if not A or not B:
        return 1.0 if a == b else 0.0
    inter = sum(min(v, B.get(k, 0)) for k, v in A.items())
    return 2 * inter / (sum(A.values()) + sum(B.values()))


# ── 3. 제출 목록 읽기 ──────────────────────────────────────────────────────

def read_names(data: bytes, filename: str) -> list[str]:
    """올린 파일(xlsx/csv/txt)에서 회사명을 모은다. 머리줄(회사명·업체명 …)은 뺀다. 첫 글 열을 쓴다."""
    names: list[str] = []
    if filename.lower().endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        # ① 앞 20줄에서 짧은 머리글(「회사명」·「피감사회사명」…) 칸 — 긴 제목 줄(「2025년 피외감대상회사 현황」)은 머리글로 안 본다
        head_row, col = None, None
        for i, r in enumerate(rows[:20]):
            for k, v in enumerate(r):
                t = re.sub(r"\s+", "", str(v or ""))
                if t and len(t) <= 12 and re.search(r"회사명|업체명|법인명|상호|피감사|감사대상|피외감", t):
                    head_row, col = i, k
                    break
            if col is not None:
                break
        # ② 머리글이 없으면 한글·영문 이름이 가장 많은 열(순번·번호·날짜 열 제외)
        if col is None:
            score: dict[int, int] = {}
            for r in rows:
                for k, v in enumerate(r):
                    t = str(v or "").strip()
                    if re.search(r"[가-힣A-Za-z]", t) and not re.fullmatch(r"[\d\s.,\-/]+", t):
                        score[k] = score.get(k, 0) + 1
            col = max(score, key=score.get) if score else 0
            head_row = -1
        for r in rows[head_row + 1:]:
            v = r[col] if col < len(r) else None
            t = str(v).strip() if v is not None else ""
            if t and not re.fullmatch(r"[\d\s.,\-/]+", t):   # 순번·숫자만 있는 칸은 이름이 아니다
                names.append(t)
    else:
        text = data.decode("utf-8-sig", errors="replace")
        for line in text.splitlines():
            cell = re.split(r"[\t,]", line)[0].strip().strip('"')
            if cell:
                names.append(cell)
    names = [n for n in names if not re.fullmatch(r"(회사명|업체명|법인명|상호|피감사회사|감사대상회사|no\.?|번호|연번)", n, re.I)]
    return names   # 중복도 그대로 둔다 — 같은 회사를 두 번 감사했으면 두 줄이 맞다(보고서 수로 견준다)


# ── 4. 대사 ────────────────────────────────────────────────────────────────

def reconcile(dart_items: list[dict[str, Any]], submitted: list[str], *, threshold: float = 0.8,
              aliases: dict[str, str] | None = None) -> dict[str, Any]:
    """공시 감사보고서(회사+결산, 개별)와 제출 목록을 맞춘다. `aliases` = 사용자가 「같은 회사」로 등록한 {씻은 제출 이름: 씻은 공시 이름}."""
    from collections import defaultdict
    dart_by: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for e in dart_items:
        dart_by[norm_name(e["company"])].append(e)
    sub_by: dict[str, list[str]] = defaultdict(list)
    for n in submitted:
        sub_by[norm_name(n, aliases)].append(n)

    def company_row(es: list[dict[str, Any]]) -> dict[str, Any]:
        """같은 회사의 보고서 여럿(3개월 결산 등) → 한 줄: 결산기 목록, 주소는 최신 보고서, 정정 여부"""
        es = sorted(es, key=lambda x: x.get("period", ""))
        latest = es[-1]
        periods = [x.get("period", "") for x in es]
        note = ""
        if len(es) > 1:
            idx = sorted(int(p.split(".")[0]) * 12 + int(p.split(".")[1]) for p in periods if "." in p)
            gaps = {b - a for a, b in zip(idx, idx[1:])}
            span = {frozenset({3}): "3개월 단위", frozenset({6}): "반기 단위", frozenset({12}): "여러 해분"}.get(frozenset(gaps), "")
            note = f"{len(es)}건({'·'.join(periods)})" + (f" {span}" if span else "")
        return {"company": latest["company"], "indi": latest.get("indi"), "con": latest.get("con") or next((x.get("con") for x in reversed(es) if x.get("con")), None),
                "period": latest.get("period", ""), "periods": periods, "reports": len(es), "note": note, "corrected": any(x.get("corrected") for x in es)}

    matched, dart_only, sub_only = [], [], []
    for key, es in dart_by.items():
        row = company_row(es)
        names = sub_by.get(key, [])
        if names:
            extra = f"목록 {len(names)}줄" if len(names) > 1 else ""
            matched.append(row | {"submitted": names[0], "name_differs": plain_name(row["company"]) != plain_name(names[0]),
                                  "note": " · ".join(x for x in [row["note"], extra] if x)})
        else:
            dart_only.append(row)
    for key, names in sub_by.items():
        if key not in dart_by:
            sub_only.append(names[0] + (f" (제출 목록에 {len(names)}줄)" if len(names) > 1 else ""))
    # 안 맞은 것끼리 비슷한 이름 후보
    suggestions = []
    sub_only_keys = {norm_name(n): n for n in sub_only}
    def near(a: str, b: str) -> tuple[float, str]:
        """비슷한 정도와 까닭. 한쪽 이름이 다른 쪽에 통째로 들어 있으면(대창기업 ↔ 대창기업홀딩스, 하우존 ↔ 하우존대부) 비슷한 이름으로 본다 — 같다고는 안 함(사용자 지정 2026-10-08)."""
        d = similarity(a, b)
        short, long_ = sorted((a, b), key=len)
        if len(short) >= 3 and short != long_ and short in long_:
            return max(d, threshold), "이름 일부 포함"
        return d, ""
    used: set[str] = set()
    for e in dart_only:
        k = norm_name(e["company"])
        cands = sorted(((*near(k, sk), sn) for sk, sn in sub_only_keys.items() if sn not in used), key=lambda t: -t[0])
        if cands and cands[0][0] >= threshold:
            score, why, sn = cands[0]
            used.add(sn)
            suggestions.append({"company": e["company"], "submitted": sn, "score": round(score, 2), "why": why, "indi": e["indi"], "con": e["con"], "note": e.get("note", "")})
    sug_dart = {s["company"] for s in suggestions}
    sug_sub = {s["submitted"] for s in suggestions}
    return {
        "matched": matched, "dart_only": [e for e in dart_only if e["company"] not in sug_dart],
        "submitted_only": [n for n in sub_only if n not in sug_sub], "suggestions": suggestions,
        "counts": {"dart": len(dart_by), "submitted": len(sub_by), "matched": len(matched), "dart_only": len(dart_only) - len(suggestions),
                   "submitted_only": len(sub_only) - len(suggestions), "suggestions": len(suggestions),
                   "dart_reports": len(dart_items), "submitted_rows": len(submitted)},
    }


# ── 5. 엑셀 ────────────────────────────────────────────────────────────────

def build_xlsx(filer: str, start: date, end: date, grouped: dict[str, Any], recon: dict[str, Any] | None = None, recon_year: int | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="E2EFDA")
    bold = Font(bold=True)
    link = Font(color="0563C1", underline="single", size=8.5)   # B·C 열 주소는 8.5pt(사용자 지정)
    for y in grouped["years"]:
        ws = wb.create_sheet(str(y["year"]))
        ws["A1"] = f"{filer} · 전자공시 제출 감사보고서(보고서명 {y['year']}년) · 검색기간 {start:%Y.%m.%d}～{end:%Y.%m.%d}"
        ws["A1"].font = bold
        ws["A2"] = f"개별 감사보고서 {y['indi_count']}건 · 연결감사보고서 {y['con_count']}건 (정정본은 최신 것 하나로, 같은 회사라도 결산기가 다르면 따로)"
        ws["A2"].font = bold
        ws.append([])
        ws.append(["회사명", "감사보고서", "연결감사보고서", "결산", "집계"])
        for c in ("A4", "B4", "C4", "D4"):
            ws[c].font = bold
            ws[c].fill = head_fill
            ws[c].alignment = Alignment(horizontal="center")
        for e in y["items"]:
            ws.append([e["company"], e["indi"] or "", e["con"] or "", e["period"], "제외(결산기 하한 앞)" if e.get("old") else ""])
            r = ws.max_row
            for col, url in (("B", e["indi"]), ("C", e["con"])):
                if url:
                    ws[f"{col}{r}"].hyperlink = url
                    ws[f"{col}{r}"].font = link
        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 52
        ws.column_dimensions["C"].width = 52
        ws.column_dimensions["D"].width = 10
        ws.freeze_panes = "A5"
    if recon is not None:
        ws = wb.create_sheet("대사")
        c = recon["counts"]
        ws["A1"] = f"외감회사대사 — 공시 회사 {c['dart']}사(개별 감사보고서 {c.get('dart_reports', '')}건, {recon_year}) · 제출 목록 회사 {c['submitted']}사({c.get('submitted_rows', '')}줄)"
        ws["A1"].font = bold
        ws["A2"] = f"일치 {c['matched']} · 공시에만 있음(제출 목록에 없음) {c['dart_only']} · 제출 목록에만 있음(공시 없음·다른 이름) {c['submitted_only']} · 비슷한 이름(확인 필요) {c['suggestions']}"
        ws.append([])
        ws.append(["구분", "공시 회사명", "제출 목록 회사명", "감사보고서", "연결감사보고서", "확인할 것"])
        for col in range(1, 7):
            cell = ws.cell(row=4, column=col)
            cell.font = bold
            cell.fill = head_fill
        for k, s in enumerate(recon["suggestions"]):
            ws.append(["비슷한 이름(확인 필요)" if k == 0 else "", s["company"], s["submitted"], s["indi"] or "", s["con"] or "", f"유사도 {s['score']}" + (" — 같은 회사인지 확인" if k == 0 else "")])
        for k, e in enumerate(recon["dart_only"]):
            ws.append(["공시에만 있음" if k == 0 else "", e["company"], "", e["indi"] or "", e["con"] or "", ("제출 목록에 없음 — 누락·다른 이름으로 제출했는지 확인" if k == 0 else "") + (f" · {e['note']}" if e.get("note") else "")])
        for k, n in enumerate(recon["submitted_only"]):
            ws.append(["제출 목록에만 있음" if k == 0 else "", "", n, "", "", "공시가 없거나 다른 이름으로 공시 — 감사보고서 제출 여부 확인" if k == 0 else ""])
        for k, m in enumerate(recon["matched"]):
            note = " · ".join(x for x in ["회사명 표기 다름" if m.get("name_differs") else "", "정정 공시 있음" if m.get("corrected") else "", m.get("note", "")] if x)
            ws.append(["일치" if k == 0 else "", m["company"], m["submitted"], m["indi"] or "", m["con"] or "", note])
        for r in range(5, ws.max_row + 1):
            for col in ("D", "E"):
                v = ws[f"{col}{r}"].value
                if v:
                    ws[f"{col}{r}"].hyperlink = v
                    ws[f"{col}{r}"].font = link
        for col, w in zip("ABCDEF", (22, 36, 36, 50, 50, 46)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def default_period(today: date | None = None) -> tuple[date, date]:
    """화면이 기간을 안 줄 때: 전년 1월 1일 ~ 오늘(결산기 하한이 전년 1월이므로 그 결산의 보고서가 빠지지 않게)."""
    today = today or date.today()
    return date(today.year - 1, 1, 1), today


# ── 브라우저판 입구(api/v1/qc.py 의 같은 이름 경로와 같은 순서) ─────────────────

import base64 as _b64
import json as _json


def web_period(start: str = "", end: str = "") -> str:
    """_period 와 같다 → JSON [start, end] (YYYY-MM-DD)."""
    try:
        s = date.fromisoformat(start) if start else None
        e = date.fromisoformat(end) if end else None
    except ValueError:
        raise ValueError("기간은 YYYY-MM-DD 로 적어 주세요.") from None
    d0, d1 = default_period()
    s, e = s or d0, e or d1
    if s > e or (e - s).days > 1200:
        raise ValueError("접수일 기간은 최대 1,200일, 시작이 끝보다 앞이어야 합니다.")
    return _json.dumps([s.isoformat(), e.isoformat()])


def _grouped(pages: list[str], max_period: str, min_period: str) -> dict:
    if max_period and not re.fullmatch(r"\d{4}\.\d{2}", max_period):
        raise ValueError("결산 상한은 YYYY.MM 꼴입니다.")
    if min_period and not re.fullmatch(r"\d{4}\.\d{2}", min_period):
        raise ValueError("결산 하한은 YYYY.MM 꼴입니다.")
    if not min_period:
        min_period = default_min_period(max_period)
    rows: list = []
    for h in pages:
        r, _, _ = parse_search(h)
        rows += r
    return group_audits(rows, max_period, min_period)


def web_audited(pages_json: str, filer: str, start: str, end: str, max_period: str = "", min_period: str = "") -> str:
    g = _grouped(_json.loads(pages_json), max_period, min_period)
    return _json.dumps(g | {"filer": re.sub(r"\s+", "", filer), "start": start, "end": end}, ensure_ascii=False)


def web_audited_xlsx(pages_json: str, filer: str, start: str, end: str, max_period: str = "", min_period: str = "") -> bytes:
    g = _grouped(_json.loads(pages_json), max_period, min_period)
    return build_xlsx(re.sub(r"\s+", "", filer), date.fromisoformat(start), date.fromisoformat(end), g)


def web_reconcile(pages_json: str, file_b64: str, filename: str, filer: str, year: int, start: str, end: str,
                  max_period: str = "", min_period: str = "", as_xlsx: bool = False, aliases: str = ""):
    data = _b64.b64decode(file_b64)
    if len(data) > 8 * 1024 * 1024:
        raise ValueError("파일이 너무 큽니다(8MB).")
    try:
        names = read_names(data, filename or "")
    except Exception as exc:   # noqa: BLE001
        raise ValueError(f"목록 파일을 읽지 못했습니다: {exc}") from exc
    if not names:
        raise ValueError("파일에서 회사명을 찾지 못했습니다(첫 시트의 회사명 열).")
    grouped = _grouped(_json.loads(pages_json), max_period, min_period)
    years = grouped["years"]
    if not years:
        raise ValueError("그 기간에 제출된 감사보고서 공시가 없습니다.")
    if year == -1:
        pick = {"year": "전체", "items": [i for y in years for i in y["items"]]}
    else:
        pick = next((y for y in years if y["year"] == year), None) or years[0]
    items = [i for i in pick["items"] if i["indi"] and not i.get("old")]
    pairs: dict[str, str] = {}
    try:
        for pr in (_json.loads(aliases) if aliases else [])[:500]:
            sk, dk = norm_name(str(pr.get("submitted", ""))), norm_name(str(pr.get("dart", "")))
            if sk and dk:
                pairs[sk] = dk
    except (ValueError, AttributeError, TypeError):
        raise ValueError("같은 회사 짝(aliases)이 JSON 목록이 아닙니다.") from None
    result = reconcile(items, names, aliases=pairs)
    result["by_year"] = [{"year": y["year"], "indi_count": y["indi_count"], "con_count": y["con_count"]} for y in years]
    if as_xlsx:
        return build_xlsx(re.sub(r"\s+", "", filer), date.fromisoformat(start), date.fromisoformat(end), grouped, result, pick["year"])
    return _json.dumps(result | {"year": pick["year"], "years": [y["year"] for y in years], "submitted_names": len(names), "submitted_sample": names[:3],
                                 "filer": re.sub(r"\s+", "", filer), "start": start, "end": end, "max_period": max_period,
                                 "min_period": grouped.get("min_period", "")}, ensure_ascii=False)
